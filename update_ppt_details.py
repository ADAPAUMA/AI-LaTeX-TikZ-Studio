import sys
import pptx

def update_pptx(filepath):
    prs = pptx.Presentation(filepath)
    updated_count = 0

    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        text = run.text

                        # Replace old integration ID if any
                        if "26d3863c-6db9-4b25-a018-f114f6d5d4de" in text:
                            run.text = text.replace("26d3863c-6db9-4b25-a018-f114f6d5d4de", "bad150a7-6816-40db-9249-10be96c432f4")
                            updated_count += 1

                        # Update Technology Used / Solution references if applicable
                        if "IBM watsonx.ai REST API" in text:
                            run.text = text + "\n- IBM Watson Assistant WebChat & NLU Integration (Integration ID: bad150a7-6816-40db-9249-10be96c432f4, Region: au-syd)."
                            updated_count += 1

    prs.save(filepath)
    print(f"Updated {filepath}: {updated_count} replacements made.")

update_pptx("ibm_pp.pptx")
update_pptx("AICTE_IBM_BOB_Project_Submission_Template.pptx")

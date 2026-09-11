import sys
import pptx

def update_presentation(filepath):
    prs = pptx.Presentation(filepath)
    
    for slide_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                full_text = shape.text_frame.text
                
                # Slide 4: Proposed Solution
                if "Proposed Solution:" in full_text and "AI Chatbot Studio Agent" in full_text:
                    new_text = full_text.replace(
                        "1. AI Chatbot Studio Agent – Engages in natural dialogue using Streamlit st.chat_message to understand user intent",
                        "1. IBM Watson Assistant & AI Chatbot Agent – Engages in natural dialogue via IBM Watson Assistant WebChat (Integration: bad150a7-6816-40db-9249-10be96c432f4) and Streamlit AI Chat Studio to understand user intent"
                    )
                    shape.text_frame.text = new_text
                
                # Slide 18: GitHub links
                if "GitHub Public Repository:" in full_text:
                    if "IBM Watson Assistant Integration ID: bad150a7-6816-40db-9249-10be96c432f4" not in full_text:
                        shape.text_frame.text = full_text + "\n\n• IBM Watson Assistant Integration:\nIntegration ID: bad150a7-6816-40db-9249-10be96c432f4\nRegion: au-syd"

    prs.save(filepath)
    print(f"Successfully updated {filepath}")

update_presentation("ibm_pp.pptx")
update_presentation("AICTE_IBM_BOB_Project_Submission_Template.pptx")
